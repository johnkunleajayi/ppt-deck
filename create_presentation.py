from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Initialize presentation with 16:9 widescreen layout
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

# Theme Palette Definition
THEMES = {
    "sf_blue": {"bg": RGBColor(3, 27, 51), "title": RGBColor(56, 189, 248), "text": RGBColor(241, 245, 249), "tile": RGBColor(10, 37, 64), "accent": RGBColor(0, 161, 224)},
    "sf_light": {"bg": RGBColor(240, 249, 255), "title": RGBColor(2, 132, 199), "text": RGBColor(12, 74, 110), "tile": RGBColor(255, 255, 255), "accent": RGBColor(0, 161, 224)},
    "flow_teal": {"bg": RGBColor(4, 47, 46), "title": RGBColor(45, 212, 191), "text": RGBColor(240, 253, 250), "tile": RGBColor(13, 60, 58), "accent": RGBColor(20, 184, 166)},
    "purple_dark": {"bg": RGBColor(30, 27, 75), "title": RGBColor(167, 139, 250), "text": RGBColor(245, 243, 255), "tile": RGBColor(45, 38, 105), "accent": RGBColor(139, 92, 246)},
    "clean_slate": {"bg": RGBColor(248, 250, 252), "title": RGBColor(51, 65, 85), "text": RGBColor(15, 23, 42), "tile": RGBColor(255, 255, 255), "accent": RGBColor(0, 161, 224)},
}

def set_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, text, theme):
    # Left accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.6), Inches(0.12), Inches(0.6))
    bar.fill.solid()
    bar.fill.fore_color.rgb = theme["accent"]
    bar.line.fill.background()

    # Title text
    txBox = slide.shapes.add_textbox(Inches(1.05), Inches(0.5), Inches(11.5), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Poppins"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = theme["title"]

def create_tile(slide, left, top, width, height, theme, border_color=None):
    tile = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    tile.fill.solid()
    tile.fill.fore_color.rgb = theme["tile"]
    if border_color:
        tile.line.color.rgb = border_color
        tile.line.width = Pt(2)
    else:
        tile.line.fill.background()
    return tile

# --- SLIDE 1: TITLE SLIDE ---
slide1 = prs.slides.add_slide(blank_slide_layout)
set_background(slide1, THEMES["sf_blue"]["bg"])

tb = slide1.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.333), Inches(3.5))
tf = tb.text_frame
tf.word_wrap = True

p1 = tf.paragraphs[0]
p1.text = "SoftCode Salesforce Flow Masterclass"
p1.font.name = "Poppins"
p1.font.size = Pt(44)
p1.font.bold = True
p1.font.color.rgb = RGBColor(255, 255, 255)
p1.alignment = PP_ALIGN.CENTER

p2 = tf.add_paragraph()
p2.text = "Class 2 — Flow Logic: Decisions, Conditions & Updating Records"
p2.font.name = "Poppins"
p2.font.size = Pt(22)
p2.font.color.rgb = THEMES["sf_blue"]["title"]
p2.alignment = PP_ALIGN.CENTER

p3 = tf.add_paragraph()
p3.text = "\nFormat: Theory + Hands-on | Duration: 1 Hour | Level: Beginner → Intermediate"
p3.font.name = "Inter"
p3.font.size = Pt(14)
p3.font.color.rgb = RGBColor(203, 213, 225)
p3.alignment = PP_ALIGN.CENTER

# --- SLIDE 2: LEARNING OBJECTIVES ---
slide2 = prs.slides.add_slide(blank_slide_layout)
theme = THEMES["sf_light"]
set_background(slide2, theme["bg"])
add_header(slide2, "1. Learning Objectives", theme)

# Col 1
create_tile(slide2, Inches(1.0), Inches(1.6), Inches(5.4), Inches(5.0), theme)
tb = slide2.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(5.0), Inches(4.5))
tf = tb.text_frame
tf.word_wrap = True
tf.paragraphs[0].text = "Theoretical Competencies"
tf.paragraphs[0].font.size = Pt(20)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = theme["title"]

bullets_1 = [
    "Explain why business decision logic is essential in Flow.",
    "Master the distinction between Start Conditions & Decisions.",
    "Leverage the $Record global variable dynamically.",
    "Diagram business logic prior to opening Flow Builder."
]
for b in bullets_1:
    p = tf.add_paragraph()
    p.text = f"• {b}"
    p.font.size = Pt(15)
    p.font.color.rgb = theme["text"]

# Col 2
create_tile(slide2, Inches(6.9), Inches(1.6), Inches(5.4), Inches(5.0), theme)
tb = slide2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.0), Inches(4.5))
tf = tb.text_frame
tf.word_wrap = True
tf.paragraphs[0].text = "Practical Flow Builder Skills"
tf.paragraphs[0].font.size = Pt(20)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = theme["title"]

bullets_2 = [
    "Configure multi-branch outcomes inside Decision Elements.",
    "Perform conditional same-record field updates.",
    "Evaluate null/existential checks (e.g., Email Is Null).",
    "Test each branch path systematically."
]
for b in bullets_2:
    p = tf.add_paragraph()
    p.text = f"• {b}"
    p.font.size = Pt(15)
    p.font.color.rgb = theme["text"]

# Save the PowerPoint output
prs.save("SoftCode_Salesforce_Flow_Masterclass_Class2.pptx")